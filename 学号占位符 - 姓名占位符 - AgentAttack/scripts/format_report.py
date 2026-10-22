from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
doc=Document(R/'report.docx')
for sec in doc.sections:
 sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
 sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
 h=sec.header.paragraphs[0];h.text='软件安全课程项目 2026  |  智能体安全攻防与评估';h.style=doc.styles['Normal']
 for run in h.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('64748B')
 f=sec.footer.paragraphs[0];f.alignment=2
 run=f.add_run('学号占位符 · 姓名占位符  |  ');run.font.size=Pt(8)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
for style in doc.styles:
 if style.type==1:
  style.font.name='Noto Serif CJK SC';style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
  style.font.size=Pt(10.5);style.paragraph_format.line_spacing=1.18;style.paragraph_format.space_after=Pt(5)
  style.paragraph_format.widow_control=True
style_lookup={s.name:s for s in doc.styles}
for name,size in [('Title',22),('Subtitle',13),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
 s=style_lookup[name];s.font.name='Noto Sans CJK SC';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('193C58');s.paragraph_format.keep_with_next=True
first=True
for p in doc.paragraphs:
 if p.style.name=='Heading 1':
  p.paragraph_format.page_break_before=False;first=False
  p.paragraph_format.space_before=Pt(15);p.paragraph_format.space_after=Pt(10)
 if p.style.name in ['Title','Subtitle','Author','Date']:
  p.alignment=1
 if p.style.name=='Title':p.paragraph_format.space_before=Pt(18)
 # Permit long evidence paths and identifiers to wrap without narrowing entire tables.
 for run in p.runs:
  if 'provenance/' in run.text or 'results/' in run.text or 'patterns/' in run.text:
   run.text=run.text.replace('/','/\u200b')
for table_index,table in enumerate(doc.tables):
 table.autofit=False
 widths=[[2.1,4.95,4.95,5.0],[1.1,4.2,5.0,6.7],[1.3,2.2,6.75,6.75],[1.1,8.8,7.1]][table_index]
 for col,width in zip(table.columns,widths):col.width=Cm(width)
 for row_i,row in enumerate(table.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if row_i==0:pr.append(OxmlElement('w:tblHeader'))
  for col_i,cell in enumerate(row.cells):
   cell.width=Cm(widths[col_i]);cell.vertical_alignment=1
   if row_i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'193C58');cell._tc.get_or_add_tcPr().append(shade)
   elif row_i%2:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F4F7');cell._tc.get_or_add_tcPr().append(shade)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.12
    for run in p.runs:
     run.font.size=Pt(9)
     if row_i==0:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
     run.text=run.text.replace('/','/\u200b')
# Set Word metadata without personal identity.
doc.core_properties.author='姓名占位符';doc.core_properties.title='虚假 Skill 前置条件扩大任务范围：三个本地模拟实例';doc.core_properties.subject='2026 软件安全课程项目；证据驱动的智能体安全评估'
doc.save(R/'report.docx')
